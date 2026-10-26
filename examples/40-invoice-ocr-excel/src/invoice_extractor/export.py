"""Write extracted invoices to an Excel sheet (or CSV)."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .parse import InvoiceData

COLUMNS = ["source", "vendor", "invoice_number", "date", "due_date", "subtotal", "tax", "total",
           "currency", "problems"]
WARN = PatternFill("solid", start_color="FFF2CC")


def write_excel(invoices: list[InvoiceData], path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append([c.replace("_", " ").title() for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for inv in invoices:
        row = inv.to_dict()
        row["problems"] = "; ".join(row["problems"])
        ws.append([row[c] for c in COLUMNS])
        if row["problems"]:  # highlight rows a person should check
            for cell in ws[ws.max_row]:
                cell.fill = WARN
    for col in ("D", "E"):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd"
    for col in ("F", "G", "H"):
        for cell in ws[col][1:]:
            cell.number_format = "#,##0.00"
    for col, width in zip("ABCDEFGHIJ", (28, 28, 18, 12, 12, 12, 12, 12, 9, 40), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return Path(path)


def write_csv(invoices: list[InvoiceData], path: str | Path) -> Path:
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        for inv in invoices:
            row = inv.to_dict()
            row["problems"] = "; ".join(row["problems"])
            writer.writerow(row)
    return Path(path)
