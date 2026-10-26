from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from invoice_extractor import ocr
from invoice_extractor.cli import run
from invoice_extractor.export import write_csv
from invoice_extractor.parse import parse_amount, parse_date, parse_invoice

from .pdf_helper import make_text_pdf

SAMPLES = Path(__file__).parent.parent / "samples"


def test_us_style_invoice():
    inv = parse_invoice((SAMPLES / "example-corp.txt").read_text(), "example-corp.txt")
    assert inv.vendor == "Example Corp" and inv.invoice_number == "INV-10442"
    assert inv.date == date(2026, 10, 3) and inv.due_date == date(2026, 11, 2)
    assert (inv.subtotal, inv.tax, inv.total, inv.currency) == (1440.0, 115.2, 1555.2, "USD")
    assert inv.problems == []


def test_eu_style_receipt():
    inv = parse_invoice((SAMPLES / "sample-bakery-receipt.txt").read_text())
    assert inv.invoice_number == "2026-0815" and inv.date == date(2026, 10, 26)
    assert (inv.subtotal, inv.tax, inv.total, inv.currency) == (3.36, 0.64, 4.0, "EUR")


def test_us_dates_option():
    inv = parse_invoice((SAMPLES / "example-corp.txt").read_text(), day_first=False)
    assert inv.date == date(2026, 3, 10)


def test_problems_are_reported():
    inv = parse_invoice("Some Shop\nSubtotal 10.00\nTax 1.00\nTotal 12.00\nDue date: 01/01/2020\nDate: 05/01/2020")
    assert "invoice_number not found" in inv.problems
    assert "subtotal + tax does not match total" in inv.problems
    assert "due date before invoice date" in inv.problems


@pytest.mark.parametrize("raw,value", [("1,234.56", 1234.56), ("1.234,56", 1234.56), ("4,00", 4.0),
                                       ("$ 99.90", 99.9), ("1 200,00", 1200.0)])
def test_parse_amount(raw, value):
    assert parse_amount(raw) == value


def test_parse_date_formats():
    assert parse_date("2026-10-03") == date(2026, 10, 3)
    assert parse_date("03.10.26") == date(2026, 10, 3)
    assert parse_date("31/31/2026") is None


def test_image_ocr_is_mocked_and_preprocessed(tmp_path, monkeypatch):
    img = tmp_path / "scan.jpg"
    Image.new("RGB", (400, 300), "white").save(img)
    seen = {}

    def fake(image, lang, config):
        seen.update(size=image.size, mode=image.mode, config=config)
        return (SAMPLES / "example-corp.txt").read_text()

    monkeypatch.setattr(ocr.shutil, "which", lambda cmd: "/usr/bin/tesseract")
    monkeypatch.setattr(ocr.pytesseract, "image_to_string", fake)
    assert "INV-10442" in ocr.load_text(img)
    assert seen == {"size": (800, 600), "mode": "L", "config": "--psm 6"}


def test_folder_to_excel(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "a.txt").write_text((SAMPLES / "example-corp.txt").read_text())
    (folder / "b.pdf").write_bytes(make_text_pdf((SAMPLES / "sample-bakery-receipt.txt").read_text()))
    (folder / "c.pdf").write_bytes(make_text_pdf(""))  # no text layer
    (folder / "notes.docx").write_bytes(b"ignored")
    out = tmp_path / "invoices.xlsx"
    results = run([folder], out)
    assert [r.source for r in results] == ["a.txt", "b.pdf", "c.pdf"]
    ws = load_workbook(out)["Invoices"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("Source", "Vendor", "Invoice Number")
    assert rows[1][7] == 1555.2 and rows[1][3].date() == date(2026, 10, 3)
    assert rows[2][2] == "2026-0815" and rows[2][7] == 4.0
    assert "no text layer" in rows[3][9]
    assert ws["A4"].fill.start_color.rgb.endswith("FFF2CC") and ws.freeze_panes == "A2"


def test_csv_export(tmp_path):
    inv = parse_invoice((SAMPLES / "example-corp.txt").read_text(), "x")
    text = write_csv([inv], tmp_path / "o.csv").read_text()
    assert text.splitlines()[0].startswith("source,vendor") and "1555.2" in text


@pytest.mark.skipif(ocr.shutil.which("tesseract") is None, reason="tesseract binary not installed")
def test_real_ocr_on_rendered_invoice(tmp_path):
    from PIL import ImageDraw, ImageFont
    img = Image.new("RGB", (1200, 700), "white")
    ImageDraw.Draw(img).multiline_text((40, 40), (SAMPLES / "example-corp.txt").read_text(), fill="black",
                                       font=ImageFont.load_default(size=28), spacing=10)
    img.save(tmp_path / "scan.png")
    inv = parse_invoice(ocr.load_text(tmp_path / "scan.png"))
    assert inv.invoice_number == "INV-10442" and inv.total == 1555.2
